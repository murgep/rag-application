from pathlib import Path

import pymupdf
import pytesseract
from docx import Document
from PIL import Image
from pypdf import PdfReader

from app.models.schemas import (
    DocumentMetadata,
    ParsedDocument,
    ParsedPage,
)


def parse_txt(file_path: str) -> ParsedDocument:
    """
    Parse a TXT file into a ParsedDocument.
    """

    path = Path(file_path)

    text = path.read_text(encoding="utf-8")

    metadata = DocumentMetadata(
        document_id=path.stem,
        filename=path.name,
        file_type="txt",
        source=str(path),
    )

    page = ParsedPage(
        page_number=1,
        text=text,
        extraction_method="text",
    )

    return ParsedDocument(
        text=text,
        metadata=metadata,
        pages=[page],
    )


def ocr_pdf_page(file_path: str, page_number: int) -> str:
    """
    Render a PDF page as an image and extract text using Tesseract OCR.

    page_number is 1-based.
    """

    pdf = pymupdf.open(file_path)

    try:
        page = pdf[page_number - 1]

        # Render the PDF page at 2x resolution
        # to improve OCR accuracy.
        matrix = pymupdf.Matrix(2, 2)

        pixmap = page.get_pixmap(matrix=matrix)

        image = Image.frombytes(
            "RGB",
            [pixmap.width, pixmap.height],
            pixmap.samples,
        )

        text = pytesseract.image_to_string(image)

        return text.strip()

    finally:
        pdf.close()


def parse_pdf(file_path: str) -> ParsedDocument:
    """
    Parse a PDF.

    Normal text pages are extracted using PyPDF.
    Pages without extractable text are processed using OCR.
    """

    path = Path(file_path)

    reader = PdfReader(file_path)

    pages = []

    for page_number, page in enumerate(
        reader.pages,
        start=1,
    ):
        # First attempt normal text extraction.
        text = page.extract_text() or ""

        if text.strip():
            extraction_method = "text"

        else:
            # No usable text found.
            # Fall back to OCR.
            text = ocr_pdf_page(
                file_path,
                page_number,
            )

            extraction_method = "ocr"

        pages.append(
            ParsedPage(
                page_number=page_number,
                text=text,
                extraction_method=extraction_method,
            )
        )

    # Combine all page text into one document-level string.
    full_text = "\n\n".join(
        page.text
        for page in pages
        if page.text.strip()
    )

    metadata = DocumentMetadata(
        document_id=path.stem,
        filename=path.name,
        file_type="pdf",
        source=str(path),
    )

    return ParsedDocument(
        text=full_text,
        metadata=metadata,
        pages=pages,
    )


def parse_docx(file_path: str) -> ParsedDocument:
    """
    Parse a DOCX file.

    Currently extracts normal paragraphs.
    """

    path = Path(file_path)

    document = Document(file_path)

    paragraphs = [
        paragraph.text
        for paragraph in document.paragraphs
        if paragraph.text.strip()
    ]

    full_text = "\n".join(paragraphs)

    metadata = DocumentMetadata(
        document_id=path.stem,
        filename=path.name,
        file_type="docx",
        source=str(path),
    )

    page = ParsedPage(
        page_number=1,
        text=full_text,
        extraction_method="text",
    )

    return ParsedDocument(
        text=full_text,
        metadata=metadata,
        pages=[page],
    )


def parse_document(file_path: str) -> ParsedDocument:
    """
    Determine the file type and route it
    to the appropriate parser.
    """

    path = Path(file_path)

    if not path.exists():
        raise FileNotFoundError(
            f"File not found: {file_path}"
        )

    if not path.is_file():
        raise ValueError(
            f"Path is not a file: {file_path}"
        )

    extension = path.suffix.lower()

    if extension == ".txt":
        return parse_txt(file_path)

    if extension == ".pdf":
        return parse_pdf(file_path)

    if extension == ".docx":
        return parse_docx(file_path)

    raise ValueError(
        f"Unsupported file type: {extension}"
    )